from __future__ import annotations

import io
import logging
import re
import shutil
import subprocess
import unicodedata
import zipfile
from dataclasses import dataclass
from functools import lru_cache
from io import StringIO
from pathlib import Path
from time import perf_counter
from typing import Callable, List, Optional, Tuple

from . import config
from .logger import get_logger

logger = get_logger(__name__)
# pdfminer logs a warning for every malformed font descriptor; in bulk runs that is pure noise.
for _noisy in ("pdfminer", "pdfminer.pdffont", "pdfminer.pdfinterp", "pdfminer.pdfpage", "PIL"):
    logging.getLogger(_noisy).setLevel(logging.ERROR)


class ParseError(Exception):
    """Raised with a short machine-readable code (e.g. 'unsupported_file_type')."""

    def __init__(self, code: str, detail: str = "") -> None:
        super().__init__(f"{code}: {detail}" if detail else code)
        self.code = code
        self.detail = detail


@dataclass
class ParseResult:
    text: str  # normalized: lowercase ASCII, one line per source line
    display_text: str  # same lines, original case (used for name extraction)
    elapsed: float
    ocr_used: bool = False
    error: Optional[str] = None  # set when no usable text could be produced


# Characters that commonly appear in resumes and should survive ASCII folding.
_UNICODE_MAP = {
    "\u2013": "-", "\u2014": "-", "\u2012": "-", "\u2212": "-", "\u2010": "-", "\u2011": "-",
    "\u2018": "'", "\u2019": "'", "\u201c": '"', "\u201d": '"', "\u00a0": " ",
    "\u2022": " ", "\u25cf": " ", "\u25aa": " ", "\u2023": " ", "\u27a2": " ", "\uf0b7": " ", "\uf0a7": " ",
    "\u00b7": " ", "\u2024": ".", "\u2026": "...",
}
_UNICODE_TABLE = str.maketrans(_UNICODE_MAP)
_LIGATURES = {"(cid:415)": "ti", "(cid:414)": "tt", "(cid:425)": "tf", "(cid:426)": "fi", "(cid:427)": "fl"}


def _fold_ascii(text: str) -> str:
    for src, dst in _LIGATURES.items():
        text = text.replace(src, dst)
    text = re.sub(r"\(cid:\d+\)", " ", text)
    text = text.translate(_UNICODE_TABLE)
    text = unicodedata.normalize("NFKD", text)
    return text.encode("ascii", "ignore").decode("ascii")


def _clean_lines(text: str) -> List[str]:
    text = _fold_ascii(text or "")
    lines = [re.sub(r"[ \t\f\v]+", " ", line).strip() for line in text.splitlines()]
    lines = [line for line in lines if line]
    if not lines:
        return []
    # Drop short lines repeated 3+ times (page headers/footers). Long lines are kept.
    freq: dict[str, int] = {}
    for line in lines:
        freq[line] = freq.get(line, 0) + 1
    return [line for line in lines if freq[line] < 3 or len(line) > 60]


def normalize_text(text: str) -> Tuple[str, str]:
    """Return (normalized_lowercase_text, display_text) preserving line structure."""
    lines = _clean_lines(text)
    display = "\n".join(lines)
    return display.lower(), display




def parse_txt(file_path: str) -> str:
    data = Path(file_path).read_bytes()
    for encoding in ("utf-8-sig", "utf-16") if data[:2] in (b"\xff\xfe", b"\xfe\xff") else ("utf-8-sig",):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1", errors="ignore")


_W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _docx_xml_text(file_path: str) -> str:
    """Extract text from every WordprocessingML part (body, text boxes, headers, footers)."""
    out: List[str] = []
    with zipfile.ZipFile(file_path) as zf:
        names = [n for n in zf.namelist() if re.match(r"word/(document|header\d*|footer\d*|footnotes|endnotes)\.xml$", n)]
        names.sort(key=lambda n: (not n.startswith("word/document"), n))
        for name in names:
            xml = zf.read(name).decode("utf-8", errors="ignore")
            # paragraph boundaries -> newlines, tabs/breaks -> spaces, then strip tags
            xml = re.sub(r"</w:p>", "\n", xml)
            xml = re.sub(r"<w:(tab|br|cr)\b[^>]*/>", " ", xml)
            texts = re.sub(r"<[^>]+>", "", xml)
            texts = (texts.replace("&amp;", "&").replace("&lt;", "<").replace("&gt;", ">")
                     .replace("&quot;", '"').replace("&apos;", "'"))
            out.append(texts)
    return "\n".join(out)


def parse_docx(file_path: str) -> str:
    try:
        from docx import Document

        doc = Document(file_path)
        chunks = [p.text for p in doc.paragraphs if p.text]
        for table in doc.tables:
            for row in table.rows:
                cells: List[str] = []
                for cell in row.cells:
                    value = cell.text.strip()
                    if value and (not cells or cells[-1] != value):  # merged cells repeat
                        cells.append(value)
                if cells:
                    chunks.append(" | ".join(cells))
        text = "\n".join(chunks)
    except zipfile.BadZipFile as exc:
        raise ParseError("corrupt_file", "not a valid .docx archive") from exc
    except Exception as exc:  # python-docx can fail on unusual but readable files
        logger.debug("python-docx failed on %s (%s); using raw XML", Path(file_path).name, exc)
        text = ""
    # python-docx ignores text boxes/shapes/headers; fall back to raw XML when that is richer.
    if len(text.strip()) < 400:
        try:
            xml_text = _docx_xml_text(file_path)
        except zipfile.BadZipFile as exc:
            raise ParseError("corrupt_file", "not a valid .docx archive") from exc
        if len(xml_text.strip()) > len(text.strip()):
            text = xml_text
    return text


@lru_cache(maxsize=1)
def _antiword_path() -> Optional[str]:
    return shutil.which("antiword") or next(
        (p for p in (r"C:\Program Files\Git\mingw64\bin\antiword.exe", r"C:\Program Files\Git\usr\bin\antiword.exe")
         if Path(p).exists()), None)


def parse_doc(file_path: str) -> str:
    exe = _antiword_path()
    if not exe:
        raise ParseError("unsupported_file_type", ".doc needs 'antiword' installed; convert the file to .docx or .pdf")
    try:
        proc = subprocess.run([exe, "-w", "0", file_path], capture_output=True, timeout=60)
    except subprocess.TimeoutExpired as exc:
        raise ParseError("parse_timeout", "antiword timed out") from exc
    if proc.returncode != 0:
        raise ParseError("parse_error", proc.stderr.decode("utf-8", "ignore").strip()[:200] or "antiword failed")
    return proc.stdout.decode("utf-8", errors="ignore")


def parse_pdf(file_path: str, max_pages: Optional[int] = None) -> str:
    from pdfminer.converter import TextConverter
    from pdfminer.layout import LAParams
    from pdfminer.pdfinterp import PDFPageInterpreter, PDFResourceManager
    from pdfminer.pdfdocument import PDFPasswordIncorrect
    from pdfminer.psparser import PSException

    max_pages = max_pages or config.MAX_PDF_PAGES
    resource_manager = PDFResourceManager(caching=True)
    buffer = StringIO()
    converter = TextConverter(resource_manager, buffer, laparams=LAParams())
    try:
        with open(file_path, "rb") as fp:
            interpreter = PDFPageInterpreter(resource_manager, converter)
            # check_extractable=False: many resumes set the "no copy" flag; we only read text.
            pages = _load_pages(fp, max_pages)
            if not pages:
                raise ParseError("corrupt_file", "unreadable PDF (no pages could be located)")
            for page in pages:
                try:
                    interpreter.process_page(page)
                except Exception as page_exc:
                    logger.debug("Skipping malformed PDF page in %s: %s", Path(file_path).name, page_exc)
    except PDFPasswordIncorrect as exc:
        raise ParseError("encrypted_pdf", "PDF is password protected") from exc
    except PSException as exc:
        raise ParseError("corrupt_file", f"unreadable PDF ({type(exc).__name__})") from exc
    finally:
        converter.close()
    text = buffer.getvalue()
    buffer.close()
    return text


def _load_pages(fp, max_pages: int) -> list:
    """PDF pages via pdfminer's page tree, falling back to _recover_pages for broken trailers."""
    from pdfminer.pdfdocument import PDFPasswordIncorrect
    from pdfminer.pdfpage import PDFPage
    from pdfminer.psparser import PSException

    try:
        pages = list(PDFPage.get_pages(fp, maxpages=max_pages, check_extractable=False))
    except PDFPasswordIncorrect:
        raise
    except PSException:
        pages = []
    return pages or _recover_pages(fp, max_pages)


_OBJ_HEADER = re.compile(rb"(?<![0-9])(\d{1,7})\s+(\d{1,5})\s+obj\b")


def _recover_pages(fp, max_pages: int) -> list:
    """Pages of a PDF whose trailer has no /Root (seen with incremental-update watermarks).

    pdfminer reads only the last trailer and its own line-based fallback misses objects, so we
    index every "N G obj" in the raw bytes (later definitions win, as in incremental updates),
    locate the /Catalog object and walk the page tree from there."""
    from pdfminer.pdfdocument import PDFDocument, PDFXRef
    from pdfminer.pdfpage import PDFPage
    from pdfminer.pdfparser import PDFParser
    from pdfminer.psparser import LIT

    fp.seek(0)
    data = fp.read()
    xref = PDFXRef()
    xref.offsets = {int(m.group(1)): (None, m.start(), int(m.group(2))) for m in _OBJ_HEADER.finditer(data)}
    xref.trailer = {}
    fp.seek(0)
    # Built by hand: PDFDocument.__init__ raises when the trailer has no /Root, which is the
    # very case being recovered. These are the attributes its object lookup relies on.
    doc = PDFDocument.__new__(PDFDocument)
    doc.caching, doc.xrefs, doc.info, doc.catalog = True, [xref], [], {}
    doc.encryption, doc.decipher = None, None
    doc._cached_objs, doc._parsed_objs = {}, {}
    doc.is_printable = doc.is_modifiable = doc.is_extractable = True
    doc._parser = PDFParser(fp)
    doc._parser.set_document(doc)
    catalog = None
    for objid in xref.offsets:
        try:
            obj = doc.getobj(objid)
        except Exception:
            continue
        if isinstance(obj, dict) and obj.get("Type") is LIT("Catalog"):
            catalog = obj
    if catalog is None:
        return []
    doc.catalog = catalog
    try:
        return list(PDFPage.create_pages(doc))[:max_pages]
    except Exception:
        return []


def _vector_glyph_pdf(file_path: str) -> bool:
    """True when pages draw lots of path segments but contain no text operators at all,
    i.e. the text was converted to vector outlines (e.g. some 'Print to PDF' outputs)."""
    from pdfminer.pdftypes import stream_value

    paths = text_ops = 0
    try:
        with open(file_path, "rb") as fp:
            for page in _load_pages(fp, 3):
                for content in page.contents or []:
                    data = stream_value(content).get_data()
                    text_ops += len(re.findall(rb"\bBT\b", data))
                    paths += len(re.findall(rb" [cl]\b", data))
    except Exception:
        return False
    return text_ops == 0 and paths > 2000


# ----------------------------------------------------------------------------- OCR


@lru_cache(maxsize=1)
def tesseract_cmd() -> Optional[str]:
    candidates = [config.TESSERACT_CMD, shutil.which("tesseract") or "",
                  r"C:\Program Files\Tesseract-OCR\tesseract.exe", r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe"]
    return next((c for c in candidates if c and Path(c).exists()), None)


def ocr_available() -> bool:
    if not config.OCR_ENABLED or not tesseract_cmd():
        return False
    try:
        import pytesseract  # noqa: F401
        from PIL import Image  # noqa: F401
    except ImportError:
        return False
    return True


def _pdf_page_images(file_path: str, max_pages: int):
    """Yield PIL images embedded in the first pages of a PDF (scanned resumes are page images).

    Uses pdfminer's own stream decoding, so no external renderer (poppler) is required."""
    from PIL import Image
    from pdfminer.pdftypes import resolve1

    with open(file_path, "rb") as fp:
        for page in _load_pages(fp, max_pages):
            xobjects = resolve1((page.resources or {}).get("XObject")) or {}
            images = []
            for obj in xobjects.values():
                stream = resolve1(obj)
                if getattr(stream, "get", None) is None or getattr(resolve1(stream.get("Subtype")), "name", "") != "Image":
                    continue
                width, height = int(resolve1(stream.get("Width")) or 0), int(resolve1(stream.get("Height")) or 0)
                if width * height < 200 * 200:  # logos/icons
                    continue
                images.append((width * height, stream, width, height))
            for _, stream, width, height in sorted(images, key=lambda x: -x[0])[:2]:
                img = _decode_pdf_image(stream, width, height, Image)
                if img is not None:
                    yield img


def _decode_pdf_image(stream, width, height, Image):
    from pdfminer.pdftypes import resolve1

    filters = resolve1(stream.get("Filter"))
    names = [getattr(resolve1(f), "name", str(f)) for f in (filters if isinstance(filters, list) else [filters]) if f]
    try:
        if "DCTDecode" in names or "JPXDecode" in names:
            # pdfminer applies any preceding filters (e.g. Flate) and leaves the JPEG bytes as-is.
            return Image.open(io.BytesIO(stream.get_data())).convert("L")
        data = stream.get_data()
        bpc = int(resolve1(stream.get("BitsPerComponent")) or 8)
        colorspace = resolve1(stream.get("ColorSpace"))
        cs_name = getattr(colorspace, "name", "") if not isinstance(colorspace, list) else getattr(resolve1(colorspace[0]), "name", "")
        mode = {"DeviceRGB": "RGB", "DeviceGray": "L", "DeviceCMYK": "CMYK"}.get(cs_name)
        if bpc == 1:
            mode = "1"
        if mode is None:
            channels = len(data) // max(1, width * height)
            mode = {1: "L", 3: "RGB", 4: "CMYK"}.get(channels)
        if mode is None:
            return None
        img = Image.frombytes(mode, (width, height), data)
        return img.convert("L")
    except Exception as exc:
        logger.debug("Could not decode PDF image (%s): %s", names, exc)
        return None


def ocr_pdf(file_path: str) -> str:
    import pytesseract

    pytesseract.pytesseract.tesseract_cmd = tesseract_cmd()
    texts = []
    for img in _pdf_page_images(file_path, config.OCR_MAX_PAGES):
        if img.width < 1200:  # upscale small scans; Tesseract wants ~300dpi text
            scale = 1200 / img.width
            img = img.resize((int(img.width * scale), int(img.height * scale)))
        texts.append(pytesseract.image_to_string(img, config="--psm 3", timeout=60))
    return "\n".join(texts)


def ocr_docx_images(file_path: str) -> str:
    """OCR the largest images embedded in a .docx (resumes saved as pictures inside Word)."""
    import pytesseract
    from PIL import Image

    pytesseract.pytesseract.tesseract_cmd = tesseract_cmd()
    texts = []
    with zipfile.ZipFile(file_path) as zf:
        media = [i for i in zf.infolist() if i.filename.startswith("word/media/")
                 and i.filename.lower().endswith((".png", ".jpg", ".jpeg", ".bmp", ".tif", ".tiff", ".gif"))]
        for info in sorted(media, key=lambda i: -i.file_size)[: config.OCR_MAX_PAGES]:
            try:
                img = Image.open(io.BytesIO(zf.read(info))).convert("L")
            except Exception:
                continue
            if img.width * img.height < 300 * 300:  # logos / icons
                continue
            if img.width < 1200:
                scale = 1200 / img.width
                img = img.resize((int(img.width * scale), int(img.height * scale)))
            texts.append(pytesseract.image_to_string(img, config="--psm 3", timeout=60))
    return "\n".join(texts)


# ----------------------------------------------------------------------------- entry point

EXT_ROUTER: dict[str, Callable[[str], str]] = {
    ".pdf": parse_pdf,
    ".docx": parse_docx,
    ".txt": parse_txt,
    ".doc": parse_doc,
}


def parse_document(file_path: str) -> ParseResult:
    """Parse a resume file. Never raises: failures are reported in ParseResult.error."""
    start = perf_counter()
    suffix = Path(file_path).suffix.lower()
    try:
        parser_fn = EXT_ROUTER.get(suffix)
        if parser_fn is None:
            raise ParseError("unsupported_file_type", f"'{suffix or 'no extension'}' is not supported")
        if Path(file_path).stat().st_size == 0:
            raise ParseError("empty_file", "file is empty")
        raw = parser_fn(file_path)
        text, display = normalize_text(raw)
        ocr_used = False
        if suffix == ".docx" and len(re.sub(r"\W", "", text)) < config.OCR_MIN_TEXT_CHARS and ocr_available():
            ocr_text, ocr_display = normalize_text(ocr_docx_images(file_path))  # scanned pages pasted into Word
            if len(ocr_text) > len(text):
                text, display, ocr_used = ocr_text, ocr_display, True
        if suffix == ".pdf" and len(re.sub(r"\W", "", text)) < config.OCR_MIN_TEXT_CHARS:
            if ocr_available():
                ocr_text, ocr_display = normalize_text(ocr_pdf(file_path))
                if len(ocr_text) > len(text):
                    text, display, ocr_used = ocr_text, ocr_display, True
            if len(re.sub(r"\W", "", text)) < 20:
                if _vector_glyph_pdf(file_path):
                    raise ParseError("vector_text_pdf", "text is drawn as vector outlines (no text layer and no "
                                     "scannable images); re-export the resume as a normal PDF or DOCX")
                reason = "image-only PDF; OCR found no text" if ocr_available() else \
                    "image-only PDF (no text layer) and OCR is unavailable (install Tesseract)"
                raise ParseError("no_text_layer", reason)
        if not text.strip():
            raise ParseError("empty_text", "no text could be extracted")
        return ParseResult(text=text, display_text=display, elapsed=perf_counter() - start, ocr_used=ocr_used)
    except ParseError as exc:
        logger.debug("Parse failed for %s: %s", Path(file_path).name, exc)
        return ParseResult("", "", perf_counter() - start, error=str(exc))
    except Exception as exc:  # malformed file: never propagate
        logger.warning("Failed parsing %s (%s): %s", Path(file_path).name, type(exc).__name__, str(exc)[:200])
        return ParseResult("", "", perf_counter() - start, error=f"parse_error: {type(exc).__name__}: {str(exc)[:200]}")

